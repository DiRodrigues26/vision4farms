"""Gera o relatório Word com tudo o que foi implementado."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ── Estilos base ──────────────────────────────────────────
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

GREEN = RGBColor(0x2E, 0x7D, 0x32)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = GREEN

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(f' — {text}')
    else:
        p.add_run(text)

def file_table(rows):
    """rows = [(ficheiro, descrição), ...]"""
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = 'Ficheiro'
    hdr[1].text = 'Descrição'
    for r in hdr:
        for p in r.paragraphs:
            p.runs[0].bold = True
    for path, desc in rows:
        row = table.add_row().cells
        row[0].text = path
        row[1].text = desc
    doc.add_paragraph()  # espaço

# ═══════════════════════════════════════════════════════════
#  DOCUMENTO
# ═══════════════════════════════════════════════════════════

# Título
title = doc.add_heading('Vision4Farms — Relatório de Desenvolvimento', level=0)
for run in title.runs:
    run.font.color.rgb = GREEN

doc.add_paragraph('Relatório das funcionalidades implementadas durante as sessões de desenvolvimento.')
doc.add_paragraph('Data: abril 2026')
doc.add_paragraph()

# ───────────────────────────────────────────────────────────
heading('1. Upload de Análises com PDF (Solo e Foliar)')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'Implementação de upload mínimo de análises (data + nome da amostra + ficheiro PDF) '
    'para análises de solo e análises foliares (yields). O utilizador seleciona o PDF no '
    'telemóvel via file_picker e submete via multipart form data.'
)

heading('Ficheiros modificados/criados', level=2)
file_table([
    ('vision4farms_api/lands/views.py', 'SoilAnalysisCreateView — POST multipart para upload de análise de solo com PDF'),
    ('vision4farms_api/lands/urls.py', 'Rota: /lands/<land_id>/soil-analyses/'),
    ('vision4farms_api/yields_app/views.py', 'YieldAnalysisCreateView — POST multipart para upload de análise foliar'),
    ('vision4farms_api/yields_app/urls.py', 'Rota: /yields/<yield_id>/analyses/'),
    ('vision4farms_api/yields_app/models.py', 'Modelos YieldsHarvests e YieldsAnalysis (movidos de models.py legacy)'),
    ('vision4farms_app/lib/shared/widgets/add_analysis_dialog.dart', 'Dialog partilhado para criar análises (solo/foliar) com seleção de PDF'),
    ('vision4farms_app/lib/features/lands/screens/land_detail_screen.dart', 'Tab Análises com FAB e _resolveMediaUrl para abrir PDFs'),
    ('vision4farms_app/lib/features/crops/screens/crop_detail_screen.dart', 'Tab Análises com seleção de yield e FAB'),
    ('vision4farms_app/pubspec.yaml', 'Adicionado file_picker: ^8.0.0+1'),
    ('vision4farms_app/android/app/src/main/AndroidManifest.xml', 'Adicionado <queries> para HTTP/HTTPS (canLaunchUrl no Android 11+)'),
])

# ───────────────────────────────────────────────────────────
heading('2. Correção do RuntimeError (app_label)')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'O servidor Django crashava ao arrancar com "vision4farms_api.models.Users doesn\'t declare '
    'an explicit app_label". Causa: imports diretos do ficheiro legacy vision4farms_api/models.py '
    '(27 modelos sem app_label). Solução: mover modelos necessários para yields_app/models.py '
    'e atualizar todos os imports.'
)

heading('Ficheiros modificados', level=2)
file_table([
    ('vision4farms_api/yields_app/models.py', 'YieldsHarvests e YieldsAnalysis movidos para cá'),
    ('vision4farms_api/yields_app/views.py', 'Import alterado para from yields_app.models'),
    ('vision4farms_api/crops/views.py', 'Import alterado para from yields_app.models'),
    ('vision4farms_api/dashboard/views.py', 'Import corrigido + removido try/except silencioso'),
])

# ───────────────────────────────────────────────────────────
heading('3. Correção da Abertura de PDFs')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'Os PDFs não abriam no telemóvel por dois motivos: (1) URL mal construído — '
    'baseUrl termina em /api mas media está em /media/; (2) Android 11+ requer '
    'declaração <queries> no manifest para canLaunchUrl funcionar com HTTP.'
)

heading('Ficheiros modificados', level=2)
file_table([
    ('vision4farms_app/lib/features/lands/screens/land_detail_screen.dart', 'Helper _resolveMediaUrl para construir URL correto'),
    ('vision4farms_app/lib/features/crops/screens/crop_detail_screen.dart', 'Mesmo helper _resolveMediaUrl'),
    ('vision4farms_app/android/app/src/main/AndroidManifest.xml', 'Adicionadas <queries> para schemes http e https'),
])

# ───────────────────────────────────────────────────────────
heading('4. Remoção do Card Estático "Geral"')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'No ecrã de lista de explorações, existia uma secção "Geral" com dados estáticos. '
    'Removida por completo, mantendo apenas o título e a lista de farms.'
)

heading('Ficheiros modificados', level=2)
file_table([
    ('vision4farms_app/lib/features/farms/screens/farm_list_screen.dart', 'Removido widget _GeneralCard e cálculo totalArea'),
])

# ───────────────────────────────────────────────────────────
heading('5. Unificação do Layout de Criar Atividade')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'O layout de criação de atividade era diferente quando aberto pela bottom bar vs. pelo '
    'ecrã de atividades. Unificado para usar sempre o mesmo dialog (showCreateActivityDialog).'
)

heading('Ficheiros modificados', level=2)
file_table([
    ('vision4farms_app/lib/features/activities/screens/activities_screen.dart', 'Removido _NewActivityModal (~215 linhas), passa a chamar showCreateActivityDialog'),
])

# ───────────────────────────────────────────────────────────
heading('6. Desenho de Polígono no Mapa (Terrenos)')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'Ao criar um terreno, o utilizador pode agora desenhar o polígono diretamente no mapa '
    '(Mapbox satellite). Toca para adicionar vértices, com undo e clear. O polígono é guardado '
    'como GeoJSON no campo land_sketch e a área em hectares é calculada automaticamente '
    'pela fórmula Shoelace com projeção esférica.'
)

heading('Ficheiros criados/modificados', level=2)
file_table([
    ('vision4farms_app/lib/features/lands/screens/draw_polygon_screen.dart', 'NOVO — Ecrã Mapbox com tap-to-draw, GeoJsonSource, FillLayer, LineLayer, CircleLayer'),
    ('vision4farms_app/lib/features/lands/screens/create_land_dialog.dart', 'Adicionado GPS, botão "Desenhar no mapa", cálculo automático de hectares, SingleChildScrollView'),
    ('vision4farms_api/lands/serializers.py', 'Adicionado land_sketch ao LandCreateSerializer'),
])

# ───────────────────────────────────────────────────────────
heading('7. Perfil do Utilizador')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'Implementação completa do perfil: edição de dados pessoais (todos os campos da BD), '
    'upload de foto de perfil via image_picker, alteração de password, e navegação para '
    'notificações e suporte.'
)

heading('7.1 Backend — Novos endpoints', level=2)
file_table([
    ('vision4farms_api/authentication/views.py', 'ProfileUpdateView (PATCH), ProfilePictureUploadView (POST), PasswordChangeView (POST)'),
    ('vision4farms_api/authentication/urls.py', 'Rotas: /auth/profile/, /auth/profile/picture/, /auth/password-change/'),
    ('vision4farms_api/authentication/serializers.py', 'ProfileUpdateSerializer com todos os campos editáveis da BD'),
])

heading('7.2 Frontend — Ecrãs e providers', level=2)
file_table([
    ('vision4farms_app/lib/features/profile/screens/profile_screen.dart', 'Refeito — avatar com foto ou letra, 4 menu items funcionais (Editar, Password, Notificações, Suporte)'),
    ('vision4farms_app/lib/features/profile/screens/edit_profile_screen.dart', 'NOVO — Formulário com nome, email, telemóvel, telefone, NIF, NIFAP, cartão fitossanitário, foto'),
    ('vision4farms_app/lib/features/profile/screens/change_password_screen.dart', 'NOVO — Password atual + nova + confirmação, com toggle de visibilidade'),
    ('vision4farms_app/lib/core/providers/auth_provider.dart', 'Adicionados updateProfile(), uploadProfilePicture(), changePassword()'),
    ('vision4farms_app/lib/core/models/user_model.dart', 'ProfileModel com profilePhone, profileNif, profileNifap, profileCardfit'),
    ('vision4farms_app/lib/core/constants/app_constants.dart', 'Endpoints profileUpdate, profilePicture, passwordChange'),
])

heading('7.3 Campos editáveis do perfil (tabela profiles)', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdr[0].text = 'Campo BD'
hdr[1].text = 'Label'
hdr[2].text = 'Tipo'
for col_name, label, tipo in [
    ('profile_name', 'Nome', 'Texto'),
    ('profile_email', 'Email', 'Email'),
    ('profile_mobile', 'Telemóvel', 'Telefone'),
    ('profile_phone', 'Telefone', 'Telefone'),
    ('profile_nif', 'NIF', 'Número (9 dígitos)'),
    ('profile_nifap', 'NIFAP', 'Número (9 dígitos)'),
    ('profile_cardfit', 'Cartão Fitossanitário', 'Número (9 dígitos)'),
    ('profile_picture', 'Foto de perfil', 'Imagem (JPG/PNG/WebP)'),
]:
    row = table.add_row().cells
    row[0].text = col_name
    row[1].text = label
    row[2].text = tipo
doc.add_paragraph()

# ───────────────────────────────────────────────────────────
heading('8. Sistema de Notificações')
# ───────────────────────────────────────────────────────────

doc.add_paragraph(
    'Sistema de notificações completo com criação automática, dois modos de visualização '
    '(overlay no dashboard e ecrã completo no perfil), toggle lido/não lido, detalhe em '
    'bottom sheet, e lembretes automáticos 24h antes dos eventos.'
)

heading('8.1 Criação automática de notificações', level=2)
doc.add_paragraph('Notificações são criadas automaticamente nos seguintes momentos:')
bullet('Ao criar uma atividade → "Nova atividade agendada" + data', bold_prefix='Atividade')
bullet('Ao criar um evento na agenda → "Novo evento na agenda" + data (já existia)', bold_prefix='Agenda')
bullet('Ao criar uma observação → "Nova observação registada" + terreno', bold_prefix='Observação')
bullet('24h antes de atividades/agenda pendentes → "Lembrete: nome" (management command)', bold_prefix='Lembrete 24h')

heading('8.2 Dois modos de visualização', level=2)
bullet('Sino na dashboard → showNotificationsOverlay() → dialog/overlay com lista', bold_prefix='Overlay')
bullet('Perfil → Notificações → openNotificationsPage() → ecrã completo com AppBar e pull-to-refresh', bold_prefix='Ecrã completo')
doc.add_paragraph('Ao clicar numa notificação individual (em ambos os modos), abre um bottom sheet com o detalhe completo.')

heading('8.3 Toggle lido/não lido', level=2)
doc.add_paragraph(
    'Cada notificação tem um ícone para alternar entre lida e não lida. '
    'Novo endpoint: PATCH /api/notifications/<id>/toggle-read/. '
    'Update otimista no frontend com rollback em caso de erro.'
)

heading('8.4 Lembrete 24h — Management command', level=2)
doc.add_paragraph(
    'O comando python manage.py send_reminders verifica atividades pendentes e '
    'eventos de agenda nas próximas 24 horas e cria notificações de lembrete. '
    'Evita duplicados verificando se já existe lembrete para o mesmo item. '
    'Recomendado agendar via cron a cada hora.'
)

heading('8.5 Ficheiros modificados/criados', level=2)
file_table([
    ('vision4farms_api/notifications_app/views.py', 'NotificationToggleReadView + agenda_id/observation_id no serializer'),
    ('vision4farms_api/notifications_app/urls.py', 'Rota: /notifications/<id>/toggle-read/'),
    ('vision4farms_api/notifications_app/management/commands/send_reminders.py', 'NOVO — Management command para lembretes 24h'),
    ('vision4farms_api/activities/serializers.py', 'Notificação ao criar atividade e observação'),
    ('vision4farms_api/agenda/serializers.py', 'Notificação ao criar evento (já existia)'),
    ('vision4farms_app/lib/features/notifications/screens/notifications_screen.dart', 'Refeito — overlay (dashboard) + ecrã completo (perfil) + detalhe em bottom sheet'),
    ('vision4farms_app/lib/features/profile/screens/profile_screen.dart', 'Usa openNotificationsPage() para ecrã completo'),
    ('vision4farms_app/lib/features/dashboard/screens/dashboard_screen.dart', 'Mantém showNotificationsOverlay() para overlay'),
    ('vision4farms_app/lib/core/constants/app_constants.dart', 'Endpoint notificationToggleRead'),
])

# ───────────────────────────────────────────────────────────
heading('9. Resumo de Todos os Ficheiros Alterados')
# ───────────────────────────────────────────────────────────

heading('Backend (Django)', level=2)
backend_files = [
    'vision4farms_api/authentication/views.py',
    'vision4farms_api/authentication/urls.py',
    'vision4farms_api/authentication/serializers.py',
    'vision4farms_api/activities/serializers.py',
    'vision4farms_api/activities/views.py',
    'vision4farms_api/agenda/serializers.py',
    'vision4farms_api/lands/views.py',
    'vision4farms_api/lands/urls.py',
    'vision4farms_api/lands/serializers.py',
    'vision4farms_api/yields_app/views.py',
    'vision4farms_api/yields_app/urls.py',
    'vision4farms_api/yields_app/models.py',
    'vision4farms_api/crops/views.py',
    'vision4farms_api/dashboard/views.py',
    'vision4farms_api/notifications_app/views.py',
    'vision4farms_api/notifications_app/urls.py',
    'vision4farms_api/notifications_app/management/commands/send_reminders.py (NOVO)',
]
for f in backend_files:
    doc.add_paragraph(f, style='List Bullet')

heading('Frontend (Flutter)', level=2)
frontend_files = [
    'vision4farms_app/lib/core/constants/app_constants.dart',
    'vision4farms_app/lib/core/models/user_model.dart',
    'vision4farms_app/lib/core/providers/auth_provider.dart',
    'vision4farms_app/lib/features/profile/screens/profile_screen.dart',
    'vision4farms_app/lib/features/profile/screens/edit_profile_screen.dart (NOVO)',
    'vision4farms_app/lib/features/profile/screens/change_password_screen.dart (NOVO)',
    'vision4farms_app/lib/features/notifications/screens/notifications_screen.dart',
    'vision4farms_app/lib/features/activities/screens/activities_screen.dart',
    'vision4farms_app/lib/features/lands/screens/create_land_dialog.dart',
    'vision4farms_app/lib/features/lands/screens/draw_polygon_screen.dart (NOVO)',
    'vision4farms_app/lib/features/lands/screens/land_detail_screen.dart',
    'vision4farms_app/lib/features/crops/screens/crop_detail_screen.dart',
    'vision4farms_app/lib/features/farms/screens/farm_list_screen.dart',
    'vision4farms_app/lib/shared/widgets/add_analysis_dialog.dart (NOVO)',
    'vision4farms_app/pubspec.yaml',
    'vision4farms_app/android/app/src/main/AndroidManifest.xml',
]
for f in frontend_files:
    doc.add_paragraph(f, style='List Bullet')

# ── Guardar ───────────────────────────────────────────────
output = r'c:\Users\diogo\Desktop\vision4farms_project\docs\Vision4Farms_Relatorio_Desenvolvimento.docx'
doc.save(output)
print(f'Documento criado: {output}')
